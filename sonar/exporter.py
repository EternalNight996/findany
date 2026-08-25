# -*- coding: utf-8 -*-
"""输出层：Excel 明细导出（含摘要） + 命中文件落盘。"""
from __future__ import annotations

import csv
import os
import shutil
import time
import re
from dataclasses import dataclass
from typing import List, Optional

from .scanner import ScanItem, ScanSummary
from .config import SearchConfig


@dataclass
class BatchDir:
    """当前日期时间批次目录。"""
    path: str
    name: str


def make_batch_dir(out_root: str) -> BatchDir:
    """out_root 下建「YYYY-MM-DD_HH-MM」日期时间批次目录，同分钟重复自动加序号。"""
    name = time.strftime("%Y-%m-%d_%H-%M")
    base = os.path.join(out_root, name)
    path = base
    n = 2
    while os.path.exists(path):
        path = base + f"_{n}"
        n += 1
    os.makedirs(path, exist_ok=True)
    return BatchDir(path=path, name=name)


def _size_str(size: int) -> str:
    v = float(size)
    for unit in ("B", "KB", "MB", "GB"):
        if v < 1024 or unit == "GB":
            return f"{v:.1f} {unit}"
        v /= 1024
    return str(size)


# Excel 单元格不接受控制字符（\x00-\x08 等），写前清理
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _clean(s):
    if not isinstance(s, str):
        return s
    return _ILLEGAL.sub("", s)


def _cell(v):
    return _clean(v) if isinstance(v, str) else v


def _ext(name: str) -> str:
    if "." not in name:
        return ""
    return name.rsplit(".", 1)[1].lower()


def _try_openpyxl():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        return Workbook, Font, PatternFill, Alignment, get_column_letter
    except Exception:
        return None


def export_excel(path: str, items: List[ScanItem], summary: ScanSummary, cfg: SearchConfig) -> str:
    """导出 Excel 明细 + 配置摘要；openpyxl 不可用时降级为 CSV。返回实际生成文件路径。"""
    rows = [i for i in items if i.hit or cfg.record_miss]
    headers = ["序号", "相对路径", "目录", "扩展名", "包含状态",
               "命中行号", "命中行内容", "匹配计数", "大小", "修改时间", "编码", "绝对路径"]

    wb_tuple = _try_openpyxl()
    if wb_tuple is None:
        csv_path = os.path.splitext(path)[0] + ".csv"
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(headers)
            for idx, it in enumerate(rows, start=1):
                state = "命中" if it.hit else "未命中"
                lines = ",".join(str(x) for x in it.hit_lines) if it.hit else ""
                w.writerow([_cell(idx), _cell(it.rel_path), _cell(it.dir_name), _cell(_ext(it.filename)),
                            _cell(state), _cell(lines), _cell(it.hit_line_text if it.hit else ""),
                            _cell(it.hit_count if it.hit else 0), _cell(_size_str(it.size)),
                            _cell(it.mtime_str), _cell(it.encoding), _cell(it.abs_path)])
        return csv_path

    Workbook, Font, PatternFill, Alignment, get_column_letter = wb_tuple
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    ws.append(headers)
    for c in ws[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for idx, it in enumerate(rows, start=1):
        state = "命中" if it.hit else "未命中"
        lines = ",".join(str(x) for x in it.hit_lines) if it.hit else ""
        ws.append([_cell(idx), _cell(it.rel_path), _cell(it.dir_name), _cell(_ext(it.filename)),
                   _cell(state), _cell(lines), _cell(it.hit_line_text if it.hit else ""),
                   _cell(it.hit_count if it.hit else 0), _cell(_size_str(it.size)),
                   _cell(it.mtime_str), _cell(it.encoding), _cell(it.abs_path)])

    widths = [6, 46, 26, 10, 10, 16, 60, 10, 10, 20, 14, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # 摘要 sheet
    ws2 = wb.create_sheet(title="摘要")
    ws2.append(["项目", "内容"])
    ws2["A1"].font = Font(bold=True)
    mode = "包含" if cfg.mode == "inc" else "不包含"
    summary_rows = [
        ("扫描目录", cfg.root_dir),
        ("关键字", cfg.keyword),
        ("匹配模式", mode),
        ("并发线程", cfg.threads),
        ("扩展名过滤", ", ".join(cfg.extensions)),
        ("编码", cfg.encoding),
        ("大小写敏感", "是" if cfg.case_sensitive else "否"),
        ("文件总数", summary.total_files),
        ("已扫描文件", summary.scanned),
        ("命中文件", summary.hit),
        ("未命中文件", summary.miss),
        ("跳过文件", summary.skipped),
        ("耗时(秒)", summary.elapsed),
        ("导出时间", time.strftime("%Y-%m-%d %H:%M:%S")),
        ("Excel 文件", os.path.basename(path) if os.path.splitext(path)[1] == ".xlsx" else "CSV"),
    ]
    for k, v in summary_rows:
        ws2.append([k, v])
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 70

    wb.save(path)
    return path


def copy_hits(items: List[ScanItem], batch_dir: str, mode: str) -> int:
    """把命中的文件按「目录名/文件名」落盘到批次目录。返回复制数。"""
    copied = 0
    for it in items:
        if not it.hit:
            continue
        target_dir = os.path.join(batch_dir, it.dir_name.replace("/", os.sep))
        try:
            os.makedirs(target_dir, exist_ok=True)
        except Exception:
            continue
        src = it.abs_path
        dst = os.path.join(target_dir, it.filename)
        if os.path.exists(dst):
            base, ext = os.path.splitext(it.filename)
            n = 2
            while os.path.exists(dst):
                dst = os.path.join(target_dir, f"{base}_{n}{ext}")
                n += 1
        try:
            shutil.copy2(src, dst)
            copied += 1
        except Exception:
            continue
    return copied
