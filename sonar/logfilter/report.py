# -*- coding: utf-8 -*-
"""筛选批次产物：filter_result.xlsx（CSV 降级）+ upload-result.csv 审计 + 命中日志留存。

审计文件不落 SecretKey / HardwareHash / payload（delivery-sop.md 第 6 条）。
"""
from __future__ import annotations

import csv
import os
import re
import shutil
import time
from typing import Dict, List

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

# 统一模板：35 列固定集合，按 6 组逻辑排序（识别→设备→网络→OA3→原始→结果）。
# 导出时「整列全空自动隐藏」——模板统一，视图按批次自适应。
DETAIL_COLS = [
    # 识别
    ("序号", "idx"), ("文件", "log_file"), ("目录", "dir_name"), ("相对路径", "rel_path"),
    ("工位", "station"), ("判型", "detected_type"),
    # 设备
    ("SN", "sn"), ("生产编号", "production_num"), ("系统SN", "system_sn"), ("板卡SN", "board_sn"),
    ("UUID", "uuid"), ("BIOS版本", "bios_version"), ("OS激活码", "os_key"),
    # 网络
    ("有线MAC", "lan"), ("无线MAC", "wifilan"), ("蓝牙MAC", "bluetooth"),
    # OA3
    ("OA3结果", "oa3_result"), ("ProductKeyID", "product_key_id"), ("PKState", "product_key_state"),
    ("ProductKey", "product_key"), ("Hash长度", "hardware_hash_len"), ("Hash SHA-256", "hardware_hash_sha256"),
    ("注入开始", "inject_start_at"), ("注入结束", "inject_end_at"),
    ("Baseboard", "baseboard_product"), ("批次号", "mo_lot_no"), ("工位任务", "task_tag"),
    # 原始
    ("JSON状态", "json_state"), ("JSON res_value", "json_res_value"), ("项目版本", "project_version"),
    # 结果
    ("提取状态", "extract_state"),
    ("回传状态", "upload_state"), ("退出码", "upload_code"), ("request_id", "request_id"),
    ("回传错误", "upload_error"),
]

# 固定宽度偏好（未列出的按内容自适应，上限 40）
COL_WIDTHS = {
    "idx": 6, "log_file": 42, "dir_name": 10, "rel_path": 46, "detected_type": 14,
    "sn": 34, "system_sn": 34, "production_num": 34, "uuid": 38, "os_key": 30,
    "hardware_hash_sha256": 20, "product_key": 30, "request_id": 22, "upload_error": 40,
}


def _col_width(key: str, header: str, samples) -> float:
    base = COL_WIDTHS.get(key)
    if base is None:
        longest = max([len(header)] + [len(str(s)) for s in samples[:80]])
        base = min(max(longest + 2, 9), 40)
    return base
AUDIT_COLS = ["时间", "SN", "日志文件", "回传状态", "退出码", "resp_status", "request_id", "尝试次数", "耗时s", "错误"]


def _cell(v):
    if isinstance(v, str):
        return _ILLEGAL.sub("", v)
    return v


def _try_openpyxl():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        return Workbook, Font, PatternFill, Alignment, get_column_letter
    except Exception:
        return None


def export_filter_excel(path: str, rows: List[Dict], summary_text: str) -> str:
    """明细 sheet + 摘要；openpyxl 缺失时降级 CSV（utf-8-sig）。返回实际路径。"""
    wb = _try_openpyxl()
    if wb is None:
        csv_path = os.path.splitext(path)[0] + ".csv"
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow([c for c, _ in DETAIL_COLS])
            for r in rows:
                w.writerow([_cell(r.get(k, "")) for _, k in DETAIL_COLS])
        return csv_path

    Workbook, Font, PatternFill, Alignment, get_column_letter = wb
    book = Workbook()
    ws = book.active
    ws.title = "明细"
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4472C4")
    ws.append([c for c, _ in DETAIL_COLS])
    for col in ws[1]:
        col.font, col.fill = hf, hfill
        col.alignment = Alignment(horizontal="center", vertical="center")
    for i, r in enumerate(rows, start=1):
        vals = [i] + [_cell(r.get(k, "")) for _, k in DETAIL_COLS[1:]]
        ws.append(vals)
    # 统一模板 + 批次自适应：整列全空 → 隐藏；宽度 = 固定偏好或按内容自适应
    for j, (header, key) in enumerate(DETAIL_COLS, start=1):
        letter = get_column_letter(j)
        samples = [str(r.get(key, "")) for r in rows if r.get(key) not in (None, "")]
        if not samples and key != "idx":
            ws.column_dimensions[letter].hidden = True      # 本批未命中的字段整列隐藏
            continue
        ws.column_dimensions[letter].width = _col_width(key, header, samples)
        ws.column_dimensions[letter].hidden = False
    ws.freeze_panes = "C2"   # 冻结表头 + 序号/文件两列，横向滚动不迷路

    ws2 = book.create_sheet(title="摘要")
    ws2.append(["项目", "内容"])
    ws2["A1"].font = Font(bold=True)
    for k, v in summary_text:
        ws2.append([k, v])
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 80
    book.save(path)
    return path


def write_upload_audit(path: str, rows: List[Dict]) -> str:
    """回传审计 CSV：sn/状态/退出码/request_id/耗时/错误。"""
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(AUDIT_COLS)
        for r in rows:
            w.writerow([
                time.strftime("%Y-%m-%d %H:%M:%S"), r.get("sn", ""), r.get("log_file", ""),
                r.get("upload_state", ""), r.get("upload_code", ""), r.get("resp_status", ""),
                r.get("request_id", ""), r.get("upload_attempts", ""), r.get("upload_elapsed", ""),
                _cell(r.get("upload_error", "")),
            ])
    return path


def copy_logs(items: List[dict], batch_dir: str) -> int:
    """提取成功的日志按「目录/文件名」留存到批次目录，同名自动加序号。"""
    copied = 0
    for it in items:
        if not it.get("extract_ok"):
            continue
        target_dir = os.path.join(batch_dir, str(it.get("dir_name", ".")).replace("/", os.sep))
        try:
            os.makedirs(target_dir, exist_ok=True)
            dst = os.path.join(target_dir, it["filename"])
            if os.path.exists(dst):
                base, ext = os.path.splitext(it["filename"])
                n = 2
                while os.path.exists(dst):
                    dst = os.path.join(target_dir, f"{base}_{n}{ext}")
                    n += 1
            shutil.copy2(it["abs_path"], dst)
            copied += 1
        except Exception:
            continue
    return copied
